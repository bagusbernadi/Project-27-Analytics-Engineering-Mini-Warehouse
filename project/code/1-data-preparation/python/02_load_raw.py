"""Langkah 4: memuat file Excel harian ke lapisan raw di DuckDB (schema raw) tanpa mengubah isinya.

Masukan : data/raw/daily/produksi_YYYYMMDD.xlsx (3 sheet: production, defect, rework), data/raw/file_manifest.csv.
Keluaran: warehouse.duckdb dengan schema raw: production, defect, rework (semua kolom VARCHAR apa adanya + _source_file, _excel_row),
          file_columns (kolom tiap sheet tiap file, dipakai test skema), file_manifest (tanggal bisnis dan tanggal diterima).
Opsi    : --through YYYY-MM-DD  hanya memuat file dengan received_date <= tanggal itu (simulasi file datang tiap hari);
          --append              menambah file yang belum dimuat tanpa menghapus database (run harian).
Tanpa --append database dibuat dari nol.
"""
import argparse
import sys
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import common as C  # noqa: E402

ap = argparse.ArgumentParser()
ap.add_argument("--through", default=None)
ap.add_argument("--append", action="store_true")
args = ap.parse_args()

man = pd.read_csv(C.RAW_DIR / "file_manifest.csv")
if args.through:
    man = man[man.received_date <= args.through]
con = duckdb.connect(str(C.DB_PATH))
con.execute("CREATE SCHEMA IF NOT EXISTS raw")
if not args.append:
    for t in ("production", "defect", "rework", "file_columns", "file_manifest"):
        con.execute(f"DROP TABLE IF EXISTS raw.{t}")
exists = {r[0] for r in con.execute("SELECT table_name FROM information_schema.tables WHERE table_schema = 'raw'").fetchall()}
done = {r[0] for r in con.execute("SELECT file_name FROM raw.file_manifest").fetchall()} if "file_manifest" in exists else set()
todo = man[~man.file_name.isin(done)]

frames = {"production": [], "defect": [], "rework": []}
fcols = []
for r in todo.itertuples():
    wb = load_workbook(C.DAILY_DIR / r.file_name, read_only=True, data_only=True)
    for sheet in ("production", "defect", "rework"):
        rows = list(wb[sheet].iter_rows(values_only=True))
        header = [str(h) for h in rows[0]]
        fcols.append((r.file_name, sheet, "|".join(header)))
        df = pd.DataFrame(rows[1:], columns=header).astype(object)
        df = df.map(lambda v: None if v is None or (isinstance(v, float) and v != v) else str(v))
        df["_source_file"] = r.file_name
        df["_excel_row"] = range(2, len(df) + 2)
        frames[sheet].append(df)
    wb.close()

for sheet, fl in frames.items():
    if not fl:
        continue
    df = pd.concat(fl, ignore_index=True)
    if sheet in {r[0] for r in con.execute("SELECT table_name FROM information_schema.tables WHERE table_schema = 'raw'").fetchall()}:
        cols = [r[0] for r in con.execute(f"DESCRIBE raw.{sheet}").fetchall()]
        for c in df.columns:                              # kolom baru muncul (perubahan skema): tambahkan ke tabel raw
            if c not in cols:
                con.execute(f'ALTER TABLE raw.{sheet} ADD COLUMN "{c}" VARCHAR')
        cols = [r[0] for r in con.execute(f"DESCRIBE raw.{sheet}").fetchall()]
        df = df.reindex(columns=cols)
        con.register("df_in", df)
        con.execute(f"INSERT INTO raw.{sheet} SELECT * FROM df_in")
        con.unregister("df_in")
    else:
        con.register("df_in", df.astype(object))
        con.execute(f"CREATE TABLE raw.{sheet} AS SELECT * FROM df_in")
        con.unregister("df_in")
if fcols:
    fc = pd.DataFrame(fcols, columns=["file_name", "sheet", "columns"])
    con.register("fc_in", fc)
    con.execute("CREATE TABLE IF NOT EXISTS raw.file_columns AS SELECT * FROM fc_in WHERE false")
    con.execute("INSERT INTO raw.file_columns SELECT * FROM fc_in")
    con.unregister("fc_in")
if len(todo):
    con.register("m_in", todo[["file_name", "business_date", "received_date"]])
    con.execute("CREATE TABLE IF NOT EXISTS raw.file_manifest AS SELECT file_name, CAST(business_date AS DATE) AS business_date, CAST(received_date AS DATE) AS received_date FROM m_in WHERE false")
    con.execute("INSERT INTO raw.file_manifest SELECT file_name, CAST(business_date AS DATE), CAST(received_date AS DATE) FROM m_in")
    con.unregister("m_in")
cnt = {t: con.execute(f"SELECT count(*) FROM raw.{t}").fetchone()[0] for t in ("production", "defect", "rework", "file_columns", "file_manifest")}
con.close()
print(f"dimuat {len(todo)} file baru;", cnt)
