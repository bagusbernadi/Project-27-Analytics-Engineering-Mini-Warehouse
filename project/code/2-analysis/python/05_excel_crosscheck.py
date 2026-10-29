"""Langkah 9 (Q2): cek silang KPI mart dengan Excel.

Dua pembanding pada sampel 20 hari kerja (setiap hari kerja ke-7), semua stasiun, model, shift:
 A. "Rekap Excel dari data bersih": baris stg_production dan stg_rework ditulis ke lembar, KPI dihitung dengan RUMUS Excel (SUMIFS) dan dibandingkan dengan fct_kpi_daily. Harus sama persis.
 B. "Rekap Excel manual dari file mentah": SUMIFS langsung pada baris mentah (teks diabaikan, kosong dihitung 0, baris ganda ikut). Selisih dengan mart dicatat dan penyebabnya dihitung per jenis data kotor.
LibreOffice menghitung ulang rumus (bukan Microsoft Excel).
Keluaran: results/2-analysis/crosscheck.xlsx (tidak di-commit), crosscheck_excel.csv (A), crosscheck_manual_diff.csv (B), crosscheck_manual_causes.csv.
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import common as C  # noqa: E402

con = duckdb.connect(str(C.DB_PATH), read_only=True)
days = [r[0] for r in con.execute("SELECT DISTINCT date FROM fct_kpi_daily ORDER BY 1").fetchall()][::7]
dl = ", ".join(f"DATE '{d}'" for d in days)
mart = con.execute(f"SELECT station_code, model, sum(input_qty) AS input_qty, sum(ok_first_pass) AS ok, sum(ng_qty) AS ng, sum(rework_qty) AS rework, sum(recovered_qty) AS recovered, sum(scrap_qty) AS scrap, "
                   f"sum(ok_first_pass)*1.0/sum(input_qty) AS fpy, sum(ng_qty)*1.0/sum(input_qty) AS ng_rate, sum(rework_qty)*1.0/sum(input_qty) AS rework_rate, sum(scrap_qty)*1.0/sum(input_qty) AS scrap_rate "
                   f"FROM fct_kpi_daily WHERE date IN ({dl}) GROUP BY 1, 2 ORDER BY 1, 2").df()
stg = con.execute(f"SELECT date, shift, station_code, model, input_qty, ok_first_pass, ng_qty FROM stg_production WHERE date IN ({dl}) ORDER BY 1, 2, 3, 4").df()
rw = con.execute(f"SELECT date, station_code, model, qty_rework, qty_recovered FROM stg_rework WHERE date IN ({dl}) ORDER BY 1, 2, 3").df()
raw = con.execute(f"""SELECT _source_file, _excel_row, CAST(left(date,10) AS DATE) AS date, shift, station, model, coalesce(input_qty, qty_input) AS input_raw, coalesce(ok_qty, qty_ok) AS ok_raw, coalesce(ng_qty, qty_ng) AS ng_raw
                      FROM raw.production WHERE CAST(left(date,10) AS DATE) IN ({dl}) ORDER BY 1, 2""").df()
dirty = pd.read_csv(C.RAW_DIR / "ground_truth_dirty.csv")
stations = pd.read_csv(C.RES_ANALYSIS / "kpi_station.csv")[["station_code", "station_name"]]
con.close()

def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None

wb = Workbook()
w1 = wb.active; w1.title = "stg"
w1.append(["date", "shift", "station_code", "model", "input_qty", "ok_first_pass", "ng_qty"])
for r in stg.itertuples(index=False):
    w1.append([r[0], r[1], r[2], r[3], int(r[4]), int(r[5]), int(r[6])])
w2 = wb.create_sheet("rework"); w2.append(["date", "station_code", "model", "qty_rework", "qty_recovered"])
for r in rw.itertuples(index=False):
    w2.append([r[0], r[1], r[2], int(r[3]), int(r[4])])
w3 = wb.create_sheet("raw")
w3.append(["station", "model", "input_raw", "ok_raw", "ng_raw"])
for r in raw.itertuples(index=False):
    w3.append([str(r.station), r.model, num(r.input_raw), num(r.ok_raw), num(r.ng_raw)])      # teks (\"480 pcs\") dan kosong menjadi sel kosong seperti SUM Excel
nS, nW, nR = len(stg) + 1, len(rw) + 1, len(raw) + 1
wc = wb.create_sheet("check", 0)
wc.append(["Cek silang: KPI dihitung dengan RUMUS Excel dari lembar stg, rework (A) dan raw (B)"]); wc["A1"].font = Font(bold=True)
hdr = ["station_code", "model", "input", "ok", "ng", "rework", "recovered", "scrap", "fpy", "ng_rate", "rework_rate", "scrap_rate", "| manual dari raw:", "input_raw", "ok_raw", "ng_raw", "fpy_raw", "ng_rate_raw"]
wc.append(hdr)
for c in wc[2]:
    c.font = Font(bold=True)
nm = stations.set_index("station_code").station_name
for i, r in enumerate(mart.itertuples(), start=3):
    wc.append([r.station_code, r.model,
               f"=SUMIFS(stg!$E$2:$E${nS},stg!$C$2:$C${nS},A{i},stg!$D$2:$D${nS},B{i})", f"=SUMIFS(stg!$F$2:$F${nS},stg!$C$2:$C${nS},A{i},stg!$D$2:$D${nS},B{i})",
               f"=SUMIFS(stg!$G$2:$G${nS},stg!$C$2:$C${nS},A{i},stg!$D$2:$D${nS},B{i})", f"=SUMIFS(rework!$D$2:$D${nW},rework!$B$2:$B${nW},A{i},rework!$C$2:$C${nW},B{i})",
               f"=SUMIFS(rework!$E$2:$E${nW},rework!$B$2:$B${nW},A{i},rework!$C$2:$C${nW},B{i})", f"=E{i}-G{i}", f"=D{i}/C{i}", f"=E{i}/C{i}", f"=F{i}/C{i}", f"=H{i}/C{i}", "",
               f'=SUMIFS(raw!$C$2:$C${nR},raw!$A$2:$A${nR},"{nm[r.station_code]}",raw!$B$2:$B${nR},B{i})', f'=SUMIFS(raw!$D$2:$D${nR},raw!$A$2:$A${nR},"{nm[r.station_code]}",raw!$B$2:$B${nR},B{i})',
               f'=SUMIFS(raw!$E$2:$E${nR},raw!$A$2:$A${nR},"{nm[r.station_code]}",raw!$B$2:$B${nR},B{i})', f"=O{i}/N{i}", f"=P{i}/N{i}"])
tmp = Path(tempfile.mkdtemp()); src = tmp / "in" / "crosscheck.xlsx"; src.parent.mkdir()
wb.save(src)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(src)], check=True, capture_output=True, timeout=300)
shutil.copy(tmp / "out" / "crosscheck.xlsx", C.RES_ANALYSIS / "crosscheck.xlsx")
v = load_workbook(C.RES_ANALYSIS / "crosscheck.xlsx", data_only=True)["check"]
xl = [r for r in v.iter_rows(min_row=3, max_row=2 + len(mart), values_only=True)]
rows, diff = [], []
for r, m in zip(xl, mart.itertuples()):
    for name, x, py, tol in (("input", r[2], m.input_qty, 0), ("ok", r[3], m.ok, 0), ("ng", r[4], m.ng, 0), ("rework", r[5], m.rework, 0), ("recovered", r[6], m.recovered, 0), ("scrap", r[7], m.scrap, 0),
                             ("fpy", r[8], m.fpy, 1e-12), ("ng_rate", r[9], m.ng_rate, 1e-12), ("rework_rate", r[10], m.rework_rate, 1e-12), ("scrap_rate", r[11], m.scrap_rate, 1e-12)):
        rows.append((m.station_code, m.model, name, x, py, "PASS" if x is not None and abs(x - py) <= tol else "FAIL"))
    diff.append(dict(station_code=m.station_code, model=m.model, input_mart=m.input_qty, input_manual=r[13], ng_mart=m.ng, ng_manual=r[15], fpy_mart=m.fpy, fpy_manual=r[16], ng_rate_mart=m.ng_rate, ng_rate_manual=r[17]))
pd.DataFrame(rows, columns=["station_code", "model", "metric", "excel", "mart", "status"]).to_csv(C.RES_ANALYSIS / "crosscheck_excel.csv", index=False)
d = pd.DataFrame(diff)
d["fpy_diff_pts"] = (d.fpy_manual - d.fpy_mart) * 100
d["ng_rate_diff_pts"] = (d.ng_rate_manual - d.ng_rate_mart) * 100
d.to_csv(C.RES_ANALYSIS / "crosscheck_manual_diff.csv", index=False)
# penyebab selisih: jenis data kotor pada baris sampel
inx = dirty[dirty.file_name.isin(set(raw._source_file))]
cause = inx.groupby("dirty_type").size().rename("rows_in_sample").reset_index()
cause["effect_on_manual_recap"] = cause.dirty_type.map({"blank_qty": "NG kosong dihitung 0: NG dan NG rate terlalu rendah", "negative_qty": "NG negatif mengurangi total NG", "text_qty": "input bertulis teks diabaikan: input terlalu kecil, FPY terlalu tinggi",
                                                          "lowercase_station": "tidak berpengaruh (SUMIFS tidak membedakan huruf besar/kecil)", "duplicate_row": "baris ganda terhitung dua kali di input, OK, dan NG"})
cause.to_csv(C.RES_ANALYSIS / "crosscheck_manual_causes.csv", index=False)
ok = pd.read_csv(C.RES_ANALYSIS / "crosscheck_excel.csv")
print(ok.status.value_counts().to_dict(), "| selisih manual: FPY maks", round(d.fpy_diff_pts.abs().max(), 3), "poin; hari sampel", len(days))
assert (ok.status == "PASS").all()
